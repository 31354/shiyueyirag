import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from docx import Document
from docx.shared import Pt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ---------- 1. 生成 PDF ----------
pdfmetrics.registerFont(TTFont("MSYH", r"C:\Windows\Fonts\msyh.ttc"))
styles = getSampleStyleSheet()
title_style = ParagraphStyle('CnTitle', parent=styles['Title'], fontName="MSYH", fontSize=20)
h1_style = ParagraphStyle('CnH1', parent=styles['Heading1'], fontName="MSYH", fontSize=14)
body_style = ParagraphStyle('CnBody', parent=styles['BodyText'], fontName="MSYH", fontSize=11)

doc = SimpleDocTemplate(os.path.join(BASE_DIR, "sample.pdf"), pagesize=A4)
story = [
    Paragraph("RAG 解析测试文档", title_style),
    Spacer(1, 0.5 * cm),
    Paragraph("1. 背景介绍", h1_style),
    Paragraph("这是一段用于测试 PDF 解析的中文段落，包含标点、数字 123 和英文 RAG。", body_style),
    Spacer(1, 0.3 * cm),
    Paragraph("2. 解析器对比", h1_style),
]
data = [
    ["解析器", "擅长格式", "输出"],
    ["Docling", "PDF/Office/HTML", "Markdown"],
    ["Marker", "PDF/EPUB", "Markdown"],
]
table = Table(data, colWidths=[4*cm, 6*cm, 4*cm])
table.setStyle(TableStyle([
    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
    ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
    ('FONTNAME', (0, 0), (-1, -1), "MSYH"),
    ('FONTSIZE', (0, 0), (-1, -1), 10),
    ('PADDING', (0, 0), (-1, -1), 6),
]))
story.append(table)
doc.build(story)
print("已生成 sample.pdf")

# ---------- 2. 生成 Word ----------
docx = Document()
docx.add_heading("RAG 解析测试文档", level=0)

docx.add_heading("1. 背景介绍", level=1)
docx.add_paragraph("这是一段用于测试 Word 解析的中文段落，包含标点、数字 456 和英文 Word。")

docx.add_heading("2. 核心组件", level=1)
for item in ["文档解析器", "分块器", "向量数据库", "大语言模型"]:
    docx.add_paragraph(item, style="List Bullet")

docx.add_heading("3. 解析器对比", level=1)
t = docx.add_table(rows=3, cols=3)
t.style = "Table Grid"
headers = ["解析器", "擅长格式", "输出"]
rows = [
    ["Docling", "PDF/Office/HTML", "Markdown"],
    ["Marker", "PDF/EPUB", "Markdown"],
]
for i, h in enumerate(headers):
    t.rows[0].cells[i].text = h
for r, row in enumerate(rows, start=1):
    for c, val in enumerate(row):
        t.rows[r].cells[c].text = val

docx.save(os.path.join(BASE_DIR, "sample.docx"))
print("已生成 sample.docx")